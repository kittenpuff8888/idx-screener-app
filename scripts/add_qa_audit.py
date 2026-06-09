from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CHECKS = [
    ("MVWAP", "VWAP", "VWAP_TP_RUNNING_VAR_V3"),
    ("MVWAP", "Price σ", "VWAP_TP_RUNNING_VAR_V3"),
    ("Previous MVWAP", "Price σ Δ 1D", "VWAP_COMPLETED_ANCHOR_DELTA_V3"),
    ("QVWAP", "Price σ", "VWAP_TP_RUNNING_VAR_V3"),
    ("Previous QVWAP", "Price σ Δ 1D", "VWAP_COMPLETED_ANCHOR_DELTA_V3"),
    ("Previous Year VWAP", "Price σ Δ 1D", "VWAP_COMPLETED_ANCHOR_DELTA_V3"),
    ("Moving Average", "EMA 25", "EMA_ADJUST_FALSE_V1"),
    ("Moving Average", "EMA 50", "EMA_ADJUST_FALSE_V1"),
    ("Moving Average", "SMA 200", "SMA_V1"),
    ("RSI", "RSI 14", "WILDER_RSI_V1"),
    ("MACD Momentum", "MACD Line", "MACD_12_26_9_V1"),
]


def grouped_columns(ws, group_row: int, header_row: int) -> dict[tuple[str, str], int]:
    current_group = ""
    columns = {}
    for col in range(1, ws.max_column + 1):
        group = str(ws.cell(group_row, col).value or "").strip()
        if group:
            current_group = group
        header = str(ws.cell(header_row, col).value or "").strip()
        if header:
            columns[(current_group, header)] = col
    return columns


def find_column(columns: dict[tuple[str, str], int], group: str, field: str):
    group_lower = group.lower()
    for (actual_group, actual_field), col in columns.items():
        actual_lower = actual_group.lower()
        if actual_field != field:
            continue
        if group == "MVWAP" and "current mvwap" in actual_lower:
            return col
        if group == "Previous MVWAP" and "previous mvwap" in actual_lower:
            return col
        if group == "QVWAP" and "current qvwap" in actual_lower:
            return col
        if group == "Previous QVWAP" and "previous qvwap" in actual_lower:
            return col
        if group == "Previous Year VWAP" and "previous year vwap" in actual_lower:
            return col
        if group_lower in actual_lower:
            return col
    return None


def add_audit(path: Path) -> None:
    wb = load_workbook(path)
    technical = wb["IDX Technical Detail"]
    columns = grouped_columns(technical, 5, 6)
    ticker_col = next(
        (col for (group, header), col in columns.items() if header == "Ticker"),
        None,
    )
    if ticker_col is None:
        raise ValueError("Could not locate the technical detail ticker column")

    name = "QA Calculation Audit"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    headers = [
        "Ticker", "Field Group", "Field", "Value", "Source",
        "Formula Version", "Missing Reason", "QA Status", "Expected Rule",
    ]
    header_fill = PatternFill("solid", fgColor="173B35")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="334B46"),
        right=Side(style="thin", color="334B46"),
        top=Side(style="thin", color="334B46"),
        bottom=Side(style="thin", color="334B46"),
    )
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_out = 2
    for row in range(7, technical.max_row + 1):
        ticker = str(technical.cell(row, ticker_col).value or "").strip()
        if not ticker:
            continue
        for group, field, version in CHECKS:
            col = find_column(columns, group, field)
            value = technical.cell(row, col).value if col else None
            missing = value in (None, "", "-", "N/A")
            completed_delta = group in {"Previous MVWAP", "Previous QVWAP", "Previous Year VWAP"}
            status = "WARN" if missing or completed_delta else "PASS"
            reason = "Missing in source workbook" if missing else (
                "Legacy workbook value; rerun corrected backend for same-anchor delta"
                if completed_delta else ""
            )
            rule = (
                "Current and prior close scored against the same completed anchor bands"
                if completed_delta else "Value present; formula version documented"
            )
            values = [ticker, group, field, value if not missing else "N/A", "Workbook/Yahoo",
                      version, reason, status, rule]
            for col_out, item in enumerate(values, start=1):
                cell = ws.cell(row_out, col_out, item)
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            row_out += 1

    widths = [12, 24, 22, 18, 18, 34, 52, 12, 56]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(1, ws.max_row)}"
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    add_audit(args.workbook)


if __name__ == "__main__":
    main()
