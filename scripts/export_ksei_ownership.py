"""Build the site's KSEI ownership data (docs/data/ksei/) and the ticker
roster / Konglo group definitions from ONE consolidated workbook:
data_sources/ksei-data-source.xlsx.

This replaced an older multi-file layout (a separate dated .xlsx per KSEI
update under data_sources/ksei/, each processed independently) with a single
file the project owner edits in place and re-commits whenever new KSEI/IDX
data is available. `asOf` always comes from the workbook's own OWNERSHIP 1%
"Date" column, never guessed from a filename.

Sheets consumed:
  TICKER LIST           -> data_sources/idx-listed.json (the roster every
                            yfinance fetch scans) + each ticker's official
                            "Sector Index" / GICS-style "Sector" / "Industry".
  SECTORAL INDEX         -> each ticker's real, published sector-index weight
                            (previously always missing -> Sectoral indices
                            silently fell back to equal-weighting every
                            constituent; this is a genuine data upgrade).
  OWNERSHIP 1%           -> per-investor holdings, long-format (one row per
                            issuer x investor), aggregated into each issuer's
                            "investors" list.
  Ownership Metrics      -> precomputed free float / HHI / CR1 / CR3 / holder
                            count / CCS per issuer (used directly; this
                            workbook already computes them, no need to
                            re-derive from the raw holder rows).
  OWNERSHIP CLASSIFICATION -> per-issuer share counts by investor type,
                            rolled into a retail/institutional/corporate/other
                            composition split (same categorisation the old
                            import_idx_ownership.py adapter used).
  KONGLO INDEX           -> data_sources/index-definitions.json's "konglo"
                            section (group -> constituent tickers).

Sector reconciliation: "Sector Index" (the official IDX code, e.g.
IDXENERGY) is blank for a real minority of tickers even in this workbook;
"Sector" is a GICS-style label sourced from yfinance and is populated far
more often but isn't IDX's own taxonomy. Every ticker's sector is resolved
with Sector Index authoritative and Sector as fallback (via
normalize_idx_sector's alias table), so a ticker is only ever "Others" when
NEITHER source classifies it -- not simply because the official field
happened to be blank.
"""
from __future__ import annotations

import json
import math
import re
import sys
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from rebuild_backend.sector_normalization import normalize_idx_sector

SOURCE_PATH = ROOT / "data_sources" / "ksei-data-source.xlsx"
OUTPUT_DIR = ROOT / "docs" / "data" / "ksei"
IDX_LISTED_PATH = ROOT / "data_sources" / "idx-listed.json"
INDEX_DEFINITIONS_PATH = ROOT / "data_sources" / "index-definitions.json"
SCHEMA_VERSION = 4

# Recent-IPO window mirrors the field's original intent: a ticker whose real
# listing date is within ~1 quarter of asOf is still building up a first
# holder table, so downstream code treats a missing ownership row as
# expected rather than a data gap.
RECENT_IPO_DAYS = 90

# Matches import_idx_ownership.py's categorisation (OWNERSHIP CLASSIFICATION
# column names), kept identical so the retail/institutional split means the
# same thing it always has.
RETAIL_COLUMNS = ["Individual (ID)"]
INSTITUTIONAL_COLUMNS = [
    "Bank", "Private Bank", "Investment Manager", "Investment Advisors", "Brokerage Firms",
    "Hedge Fund", "Sovereign Wealth Fund", "Capital Market Supporting Institutions and Professions",
    "Mutual Funds (MF)", "Securities Company (SC)", "Pension Funds (PF)",
    "Financial Institutional (IB)", "Insurance (IS)", "Private Equity", "Venture Capital",
    "Exchange Traded Funds", "Trustee Bank",
]


def finite_number(value: Any) -> float | int | None:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(parsed):
        return None
    return int(parsed) if parsed.is_integer() else parsed


def text_value(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    out = str(value).strip()
    return out if out else fallback


def normalized_name(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().upper())


def sheet_rows(ws, header_row: int = 1) -> list[dict[str, Any]]:
    headers = [text_value(ws.cell(row=header_row, column=c).value) for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {headers[c - 1]: ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1) if headers[c - 1]}
        rows.append(row)
    return rows


def parse_id_shares(value: Any) -> int | None:
    """'1.904.883.411' (Indonesian thousands separator) -> 1904883411."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    cleaned = re.sub(r"[.\s]", "", str(value).strip())
    return int(cleaned) if cleaned.isdigit() else None


def parse_listing_date(value: Any) -> str | None:
    """'08 Jun 2021' or a real datetime -> '2021-06-08'."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d %b %Y", "%d %B %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def build_ticker_list(wb: openpyxl.Workbook) -> list[dict[str, Any]]:
    out = []
    for row in sheet_rows(wb["TICKER LIST"]):
        ticker = text_value(row.get("Ticker")).upper()
        if not ticker:
            continue
        out.append({
            "ticker": ticker,
            "companyName": text_value(row.get("Company Name")),
            "listingDate": parse_listing_date(row.get("Listing Date")),
            "shares": parse_id_shares(row.get("Shares")),
            "board": text_value(row.get("Listing Board")),
            "sectorIndexRaw": text_value(row.get("Sector Index")) or None,
            "sectorRaw": text_value(row.get("Sector")) or None,
            "industry": text_value(row.get("Industry")) or None,
        })
    return out


def build_sector_weights(wb: openpyxl.Workbook) -> dict[str, dict[str, Any]]:
    """ticker -> {sectorIndex, weight} from SECTORAL INDEX's real per-ticker
    index weight (Post-Evaluation, the effective/final figure). Column
    positions are fixed (two merged header rows, not named per-cell)."""
    ws = wb["SECTORAL INDEX"]
    out: dict[str, dict[str, Any]] = {}
    for r in range(3, ws.max_row + 1):
        sector_index = text_value(ws.cell(row=r, column=1).value)
        ticker = text_value(ws.cell(row=r, column=2).value).upper()
        if not ticker or not sector_index:
            continue
        weight = finite_number(ws.cell(row=r, column=8).value)  # Index Weight, Post-Evaluation
        out[ticker] = {"sectorIndex": sector_index, "weight": weight}
    return out


def build_konglo_groups(wb: openpyxl.Workbook) -> dict[str, list[dict[str, str]]]:
    groups: dict[str, list[dict[str, str]]] = {}
    for row in sheet_rows(wb["KONGLO INDEX"]):
        group = text_value(row.get("Conglomerate Group"))
        ticker = text_value(row.get("Ticker")).upper()
        if not group or not ticker:
            continue
        groups.setdefault(group, []).append({"ticker": ticker})
    return groups


def build_ownership_metrics(wb: openpyxl.Workbook) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in sheet_rows(wb["Ownership Metrics"]):
        ticker = text_value(row.get("Ticker")).upper()
        if not ticker:
            continue
        out[ticker] = {
            "freeFloat": finite_number(row.get("Free Float (%)")),
            "hhi": finite_number(row.get("Classic HHI")),
            "cr1": finite_number(row.get("Concentration Ratio Top 1 (CR1)")),
            "cr3": finite_number(row.get("Concentration Ratio Top 3 (CR3)")),
            "holderCount": finite_number(row.get("Holder")),
            "ccs": finite_number(row.get("CCS")),
            "ownershipType": text_value(row.get("Ownership Type"), "Unclassified"),
            "ccsCategory": text_value(row.get("CCS Category"), "Unclassified"),
        }
    return out


def build_investors(wb: openpyxl.Workbook) -> tuple[dict[str, list[dict[str, Any]]], str]:
    """ticker -> ranked investors list (long-format OWNERSHIP 1% grouped by
    ticker), plus the workbook's own as-of date (its real "Date" column,
    never guessed)."""
    ws = wb["OWNERSHIP 1%"]
    by_ticker: dict[str, list[dict[str, Any]]] = {}
    as_of: str | None = None
    for row in sheet_rows(ws):
        if as_of is None:
            d = row.get("Date")
            if isinstance(d, (datetime, date)):
                as_of = d.isoformat()[:10] if isinstance(d, datetime) else d.isoformat()
        ticker = text_value(row.get("Ticker")).upper()
        name = text_value(row.get("Investor Name"))
        if not ticker or not name:
            continue
        pct = finite_number(row.get("Percentage (%)"))
        if pct is None:
            continue
        by_ticker.setdefault(ticker, []).append({
            "name": name,
            "type": text_value(row.get("Investor Classification"), "Unclassified"),
            "percentage": pct,
            "localForeign": text_value(row.get("Local / Foreign")) or None,
            "domicile": text_value(row.get("Domicile")) or None,
        })
    for ticker, holders in by_ticker.items():
        holders.sort(key=lambda h: -h["percentage"])
        for rank, holder in enumerate(holders, start=1):
            holder["rank"] = rank
            holder["originalLine"] = f"{rank}. {holder['name']} - {holder['type']} - {holder['percentage']}%"
    if as_of is None:
        raise ValueError("OWNERSHIP 1% sheet has no usable Date column -- refusing to guess asOf")
    return by_ticker, as_of


def build_composition(wb: openpyxl.Workbook) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in sheet_rows(wb["OWNERSHIP CLASSIFICATION"]):
        ticker = text_value(row.get("Ticker")).upper()
        total = finite_number(row.get("Total Scripless"))
        if not ticker or not total:
            continue
        retail = sum(finite_number(row.get(c)) or 0 for c in RETAIL_COLUMNS)
        institutional = sum(finite_number(row.get(c)) or 0 for c in INSTITUTIONAL_COLUMNS)
        corporate = finite_number(row.get("Corporate")) or 0
        other = max(total - retail - institutional - corporate, 0)
        out[ticker] = {
            "retailPct": round(100 * retail / total, 1),
            "institutionalPct": round(100 * institutional / total, 1),
            "corporatePct": round(100 * corporate / total, 1),
            "otherPct": round(100 * other / total, 1),
            "totalScripless": int(total),
            "retailShares": int(retail),
        }
    return out


def is_recent_ipo(listing_date: str | None, as_of: str) -> bool | None:
    if not listing_date:
        return None
    try:
        listed = date.fromisoformat(listing_date)
        anchor = date.fromisoformat(as_of)
    except ValueError:
        return None
    return (anchor - listed).days <= RECENT_IPO_DAYS


def write_json(path: Path, payload: Any, pretty: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if pretty:
        text = json.dumps(payload, ensure_ascii=False, indent=2)
    else:
        text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    path.write_text(text, encoding="utf-8")


def summary(records: list[dict[str, Any]]) -> dict[str, Any]:
    def values(field: str) -> list[float]:
        return [float(item[field]) for item in records if item.get(field) is not None]

    free_float = values("freeFloat")
    hhi = values("hhi")
    ownership_types = Counter(item["ownershipType"] for item in records)
    ccs_categories = Counter(item["ccsCategory"] for item in records)
    sectors = Counter(item["sector"] or "Others" for item in records)
    return {
        "totalIssuers": len(records),
        "okIssuers": sum(item["status"] == "ok" for item in records),
        "partialIssuers": sum(item["status"] == "partial" for item in records),
        "averageFreeFloat": round(sum(free_float) / len(free_float), 2) if free_float else None,
        "averageHHI": round(sum(hhi) / len(hhi), 1) if hhi else None,
        "highConcentrationIssuers": sum((item.get("hhi") or 0) >= 2500 for item in records),
        "highCCSIssuers": sum(item["ccsCategory"].lower() == "tinggi" for item in records),
        "ownershipTypes": dict(ownership_types.most_common()),
        "ccsCategories": dict(ccs_categories.most_common()),
        "sectors": dict(sectors.most_common()),
    }


def investor_directory(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    investors: dict[str, dict[str, Any]] = {}
    for issuer in records:
        for holder in issuer["investors"]:
            key = normalized_name(holder["name"])
            item = investors.setdefault(
                key,
                {"name": holder["name"], "types": Counter(), "tickers": [], "totalPublishedPercentage": 0.0},
            )
            item["types"][holder["type"]] += 1
            item["tickers"].append({
                "ticker": issuer["ticker"],
                "companyName": issuer["companyName"],
                "sector": issuer["sector"],
                "industry": issuer["industry"],
                "percentage": holder["percentage"],
                "type": holder["type"],
                "rank": holder["rank"],
                "originalLine": holder["originalLine"],
            })
            item["totalPublishedPercentage"] += holder["percentage"]
    output = []
    for item in investors.values():
        item["types"] = dict(item["types"].most_common())
        item["tickerCount"] = len(item["tickers"])
        item["totalPublishedPercentage"] = round(item["totalPublishedPercentage"], 4)
        item["tickers"].sort(key=lambda row: (-row["percentage"], row["ticker"]))
        output.append(item)
    return sorted(output, key=lambda item: (-item["tickerCount"], item["name"]))


def compare(previous: dict[str, Any] | None, current: dict[str, Any]) -> dict[str, Any]:
    if not previous:
        return {
            "previousAsOf": None, "newTickers": [], "removedTickers": [], "changedTickers": [],
            "investorAdditions": [], "investorRemovals": [],
        }
    old = {item["ticker"]: item for item in previous["records"]}
    new = {item["ticker"]: item for item in current["records"]}
    additions: list[dict[str, Any]] = []
    removals: list[dict[str, Any]] = []
    changed: list[str] = []
    for ticker in sorted(old.keys() & new.keys()):
        old_holders = {normalized_name(item["name"]): item for item in old[ticker]["investors"]}
        new_holders = {normalized_name(item["name"]): item for item in new[ticker]["investors"]}
        for key in sorted(new_holders.keys() - old_holders.keys()):
            additions.append({"ticker": ticker, **new_holders[key]})
        for key in sorted(old_holders.keys() - new_holders.keys()):
            removals.append({"ticker": ticker, **old_holders[key]})
        numeric_changed = any(
            old[ticker].get(field) != new[ticker].get(field)
            for field in ("freeFloat", "hhi", "cr1", "cr3", "holderCount", "ccs")
        )
        holder_changed = bool(old_holders.keys() ^ new_holders.keys()) or any(
            old_holders[key]["percentage"] != new_holders[key]["percentage"]
            for key in old_holders.keys() & new_holders.keys()
        )
        if numeric_changed or holder_changed:
            changed.append(ticker)
    return {
        "previousAsOf": previous["asOf"],
        "newTickers": sorted(new.keys() - old.keys()),
        "removedTickers": sorted(old.keys() - new.keys()),
        "changedTickers": changed,
        "investorAdditions": additions,
        "investorRemovals": removals,
    }


def build_idx_listed(ticker_rows: list[dict[str, Any]], as_of: str) -> dict[str, Any]:
    records = []
    for row in ticker_rows:
        records.append({
            "ticker": row["ticker"],
            "name": row["companyName"],
            "listingDate": row["listingDate"],
            "shares": row["shares"],
            "board": row["board"],
            "isRecentIpo": is_recent_ipo(row["listingDate"], as_of),
        })
    return {
        "schemaVersion": 1,
        "source": "IDX Daftar Saham (data_sources/ksei-data-source.xlsx, TICKER LIST)",
        "asOf": as_of,
        "count": len(records),
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "records": records,
    }


def update_konglo_definitions(konglo_groups: dict[str, list[dict[str, str]]]) -> None:
    definitions = json.loads(INDEX_DEFINITIONS_PATH.read_text(encoding="utf-8"))
    definitions["konglo"] = konglo_groups
    write_json(INDEX_DEFINITIONS_PATH, definitions, pretty=True)


def build() -> dict[str, Any]:
    if not SOURCE_PATH.exists():
        raise FileNotFoundError(
            f"{SOURCE_PATH} not found -- place the KSEI data source workbook there "
            "(single consolidated file: TICKER LIST / OWNERSHIP 1% / OWNERSHIP "
            "CLASSIFICATION / SECTORAL INDEX / KONGLO INDEX / Ownership Metrics sheets)."
        )
    generated_at = datetime.now(timezone.utc).isoformat()
    wb = openpyxl.load_workbook(SOURCE_PATH, data_only=True)

    ticker_rows = build_ticker_list(wb)
    sector_weights = build_sector_weights(wb)
    metrics = build_ownership_metrics(wb)
    investors_by_ticker, as_of = build_investors(wb)
    composition = build_composition(wb)
    konglo_groups = build_konglo_groups(wb)

    records: list[dict[str, Any]] = []
    for row in ticker_rows:
        ticker = row["ticker"]
        sw = sector_weights.get(ticker)
        # Sector Index (official IDX code, from either SECTORAL INDEX -- the
        # more authoritative real index-constituent listing -- or the TICKER
        # LIST's own copy of it) wins; the GICS-style "Sector" (yfinance-
        # probed) is only a fallback for the ticker's genuinely not in any
        # official sector index yet.
        primary = (sw["sectorIndex"] if sw else None) or row["sectorIndexRaw"]
        idx_sector = normalize_idx_sector(primary, row["sectorRaw"])
        m = metrics.get(ticker, {})
        investors = investors_by_ticker.get(ticker, [])
        required_present = m.get("freeFloat") is not None and m.get("hhi") is not None and m.get("ccs") is not None
        missing_fields = [] if required_present else ["freeFloat", "hhi", "ccs"]
        records.append({
            "ticker": ticker,
            "companyName": row["companyName"],
            "sector": idx_sector,
            "idxSector": idx_sector,
            "sourceSector": row["sectorRaw"] or "Others",
            "industry": row["industry"] or "Others",
            "investors": investors,
            "freeFloat": m.get("freeFloat"),
            "hhi": m.get("hhi"),
            "cr1": m.get("cr1"),
            "cr3": m.get("cr3"),
            "holderCount": m.get("holderCount"),
            "ccs": m.get("ccs"),
            "ownershipType": m.get("ownershipType", "Unclassified"),
            "ccsCategory": m.get("ccsCategory", "Unclassified"),
            "idxSectorWeight": sw["weight"] if sw else None,
            "composition": composition.get(ticker),
            "status": "ok" if required_present else "partial",
            "missingFields": missing_fields,
            "source": "KSEI data source workbook",
            "asOf": as_of,
            "formulaVersion": "ksei-ownership-v4",
        })

    dated_path = OUTPUT_DIR / "dates" / as_of / "ownership.json"
    previous: dict[str, Any] | None = None
    latest_path = OUTPUT_DIR / "latest.json"
    if latest_path.exists():
        try:
            previous = json.loads(latest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            previous = None

    payload = {
        "schemaVersion": SCHEMA_VERSION,
        "dataType": "ksei_ownership",
        "asOf": as_of,
        "generatedAt": generated_at,
        "source": {
            "name": "KSEI data source workbook",
            "file": SOURCE_PATH.name,
            "sourceMode": "point_in_time",
        },
        "summary": summary(records),
        "records": records,
        "investorChanges": [],
        "investorDirectory": investor_directory(records),
        "schemaWarnings": [],
    }
    payload["comparison"] = compare(previous if previous and previous.get("asOf") != as_of else None, payload)
    write_json(dated_path, payload)
    write_json(latest_path, payload)

    # Manifest still enumerates every dated archive ever published (including
    # ones from the old multi-file source, which stay valid historical
    # snapshots) so docs/data/ksei/trend.json keeps a real multi-point history.
    available_dates = sorted(p.name for p in (OUTPUT_DIR / "dates").iterdir() if p.is_dir())
    manifest = {
        "schemaVersion": SCHEMA_VERSION,
        "dataType": "ksei_ownership",
        "generatedAt": generated_at,
        "latestAsOf": max(available_dates),
        "availableDates": available_dates,
        "snapshots": [
            {
                "asOf": d,
                "path": f"data/ksei/dates/{d}/ownership.json",
                "recordCount": (payload["summary"]["totalIssuers"] if d == as_of else None),
                "status": "ok" if d == as_of and not payload["summary"]["partialIssuers"] else None,
            }
            for d in available_dates
        ],
    }
    write_json(OUTPUT_DIR / "manifest.json", manifest)

    idx_listed = build_idx_listed(ticker_rows, as_of)
    write_json(IDX_LISTED_PATH, idx_listed, pretty=True)
    update_konglo_definitions(konglo_groups)

    return {
        "asOf": as_of,
        "issuers": len(records),
        "investors": len(payload["investorDirectory"]),
        "konglomerateGroups": len(konglo_groups),
        "tickers": len(ticker_rows),
        "sectoralWeighted": sum(1 for r in records if r["idxSectorWeight"] is not None),
    }


if __name__ == "__main__":
    result = build()
    print(
        f"Published KSEI ownership as of {result['asOf']}: {result['issuers']} issuers, "
        f"{result['investors']} investors, {result['konglomerateGroups']} Konglo groups, "
        f"{result['tickers']} listed tickers, {result['sectoralWeighted']} with a real "
        f"sector-index weight."
    )
