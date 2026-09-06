from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from rebuild_backend.sector_normalization import normalize_idx_sector

OUTPUT_DIR = ROOT / "Output"
DOCS_DIR = ROOT / "docs"
DATA_DIR = DOCS_DIR / "data"
DOWNLOADS_DIR = DOCS_DIR / "downloads"
SOURCE_PAYLOAD_DIR = ROOT / "data_sources" / "full-workbook"

LEGACY_SCREENER_COLUMNS = [
    "Filter",
    "Ticker",
    "Sector",
    "Price",
    "Chg %",
    "RVOL",
    "ADR & ATR (14)",
    "Zone",
    "MP Profile",
    "MA Position",
    "POI",
    "Anchor",
    "Entry",
    "Target",
    "Upside %",
    "Invalidation",
    "R/R",
]

FILTER_LABELS = {
    "A": "EMA Trend",
    "B": "Golden Cross",
    "C": "Structure Break",
    "D": "Price Level Reclaim",
    "E": "Equal-Level Breakout",
    "F": "Near VWAP",
    "G": "Structure Location",
}


def clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    return value


def number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if value in (None, "", "-"):
        return None
    text = str(value).strip().replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group())
    except ValueError:
        return None


def first_value(row: dict[str, Any], names: Iterable[str]) -> Any:
    for name in names:
        value = row.get(name)
        if value not in (None, "", "-"):
            return value
    return None


def matching_value(
    row: dict[str, Any],
    *,
    contains: Iterable[str],
    suffix: str | None = None,
) -> Any:
    required = tuple(part.lower() for part in contains)
    for key, value in row.items():
        normalized = str(key).lower()
        if all(part in normalized for part in required) and (
            suffix is None or normalized.endswith(suffix.lower())
        ):
            if value not in (None, "", "-"):
                return value
    return None


def slug_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%B %d, %Y", "%B %d %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Unsupported market date: {value!r}")


def latest_workbook() -> Path:
    files = [
        path
        for path in OUTPUT_DIR.glob("IDX Screener as of *.xlsx")
        if not path.name.startswith("~$")
    ]
    if not files:
        raise FileNotFoundError(f"No IDX Screener workbook found in {OUTPUT_DIR}")
    return max(files, key=lambda path: path.stat().st_mtime)


def find_stat(ws, label: str) -> Any:
    target = re.sub(r"\s+", " ", label).strip().lower()
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        for index, value in enumerate(cells):
            normalized = re.sub(r"\s+", " ", str(value or "")).strip().lower()
            if normalized == target:
                return cells[index + 1] if index + 1 < len(cells) else None
    return None


def find_header_row(ws, required: set[str], limit: int = 30) -> int | None:
    for row_idx in range(1, min(ws.max_row, limit) + 1):
        values = {
            str(ws.cell(row_idx, col).value or "").strip()
            for col in range(1, ws.max_column + 1)
        }
        if required.issubset(values):
            return row_idx
    return None


def unique_headers(ws, header_row: int, start_col: int = 2) -> list[tuple[int, str]]:
    raw = [
        str(ws.cell(header_row, col).value or "").strip()
        for col in range(start_col, ws.max_column + 1)
    ]
    counts = Counter(header for header in raw if header)
    current_group = ""
    groups: list[str] = []
    for col in range(start_col, ws.max_column + 1):
        group = str(ws.cell(max(1, header_row - 1), col).value or "").strip()
        if group:
            current_group = re.sub(r"\s+", " ", group)
        groups.append(current_group)

    used: Counter[str] = Counter()
    result: list[tuple[int, str]] = []
    for offset, header in enumerate(raw):
        if not header:
            continue
        used[header] += 1
        name = header
        if counts[header] > 1:
            group = groups[offset]
            name = f"{group} · {header}" if group else f"{header} {used[header]}"
        name = name.replace("\u00c2\u00b7", "\u00b7")
        result.append((start_col + offset, name))
    return result


def parse_table(ws, header_row: int | None = None) -> list[dict[str, Any]]:
    header_row = header_row or find_header_row(ws, {"Ticker"})
    if not header_row:
        return []
    headers = unique_headers(ws, header_row)
    rows: list[dict[str, Any]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        item = {
            header: clean_value(ws.cell(row_idx, col).value)
            for col, header in headers
        }
        if item.get("Ticker") not in (None, "", "-"):
            rows.append(item)
    return rows


def workbook_sheet_payload(ws) -> dict[str, Any]:
    header_row = find_header_row(ws, {"Ticker"}, limit=max(30, ws.max_row))
    if header_row:
        start_col = next(
            (col for col in range(1, ws.max_column + 1) if ws.cell(header_row, col).value not in (None, "")),
            1,
        )
        headers = unique_headers(ws, header_row, start_col=start_col)
        rows = []
        for row_idx in range(header_row + 1, ws.max_row + 1):
            item = {
                header: clean_value(ws.cell(row_idx, col).value)
                for col, header in headers
            }
            if any(value not in (None, "") for value in item.values()):
                rows.append(item)
        return {
            "headerRow": header_row,
            "columns": [header for _, header in headers],
            "rows": rows,
        }

    matrix = []
    for row in ws.iter_rows(values_only=True):
        values = [clean_value(value) for value in row]
        while values and values[-1] in (None, ""):
            values.pop()
        if values:
            matrix.append(values)
    width = max((len(row) for row in matrix), default=0)
    columns = [f"Column {index}" for index in range(1, width + 1)]
    return {
        "headerRow": None,
        "columns": columns,
        "rows": [
            {columns[index]: value for index, value in enumerate(row)}
            for row in matrix
        ],
    }


def build_field_catalog(workbook_sheets: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    catalog = []
    for sheet_name, table in workbook_sheets.items():
        for column in table.get("columns", []):
            group, separator, label = column.partition(" · ")
            catalog.append({
                "sheet": sheet_name,
                "group": group if separator else "",
                "field": label if separator else column,
                "key": column,
            })
    return catalog


def build_ticker_details(workbook_sheets: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    details: dict[str, dict[str, Any]] = defaultdict(dict)
    for sheet_name, table in workbook_sheets.items():
        for row in table.get("rows", []):
            ticker = str(row.get("Ticker") or "").strip().upper()
            if ticker:
                details[ticker][sheet_name] = row
    return dict(sorted(details.items()))


def normalized_stocks(
    technical: list[dict[str, Any]],
    fundamental: list[dict[str, Any]],
    news: list[dict[str, Any]],
    screener: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    fundamentals = {str(row.get("Ticker") or "").upper(): row for row in fundamental}
    signals: dict[str, list[dict[str, Any]]] = defaultdict(list)
    news_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in screener:
        signals[str(row.get("Ticker") or "").upper()].append(row)
    for row in news:
        news_rows[str(row.get("Ticker") or "").upper()].append(row)

    stocks: dict[str, dict[str, Any]] = {}
    for row in technical:
        ticker = str(row.get("Ticker") or "").upper()
        if not ticker:
            continue
        fund = fundamentals.get(ticker, {})
        stock_signals = signals.get(ticker, [])
        primary = stock_signals[0] if stock_signals else {}
        stocks[ticker] = {
            "ticker": ticker,
            "companyName": first_value(row, ["Emiten", "Company"]) or first_value(fund, ["Company", "Emiten"]),
            "sector": normalize_idx_sector(
                first_value(row, ["IDX Sector"]),
                first_value(row, ["Sector"]) or first_value(fund, ["IDX Sector", "Sector"]),
            ),
            "industry": first_value(row, ["Industry"]) or first_value(fund, ["Industry"]),
            "lastPrice": first_value(row, ["Closing Price", "Price"]),
            "changePercent": first_value(row, ["Price Change %", "Chg %"]),
            "volume": first_value(row, ["Volume"]),
            "rvol": first_value(row, ["RVOL 20 D", "RVOL"]),
            "rvolChangePercent": first_value(row, ["RVOL Change %", "RVOL 20 D Change %"]),
            "rsRating": first_value(row, ["RS Rating"]),
            "beta": first_value(row, ["Beta (vs IHSG)", "Beta vs IHSG"]),
            "liquidityCategory": first_value(row, ["Liquidity Categories", "Liquidity Category"]),
            "trend": {
                "internal": first_value(row, ["Internal Trend"]),
                "swing": first_value(row, ["Swing Trend"]),
            },
            "structure": {
                "internal": first_value(row, ["Latest Internal Struct"]),
                "swing": first_value(row, ["Latest Swing Struct"]),
            },
            "movingAverages": {
                "ema25": first_value(row, ["EMA 25"]),
                "ema50": first_value(row, ["EMA 50"]),
                "sma200": first_value(row, ["SMA 200"]),
                "zone": first_value(row, ["MA Zone"]),
            },
            "supportLevels": [
                first_value(row, ["Strong Low"]),
                first_value(row, ["Weak Low"]),
                first_value(row, ["IBL"]),
                first_value(row, ["PWL"]),
                first_value(row, ["MDL"]),
            ],
            "resistanceLevels": [
                first_value(row, ["Strong High"]),
                first_value(row, ["Weak High"]),
                first_value(row, ["IBH"]),
                first_value(row, ["PWH"]),
                first_value(row, ["MDH"]),
            ],
            "entry": first_value(primary, ["Entry"]),
            "entryPoi": first_value(primary, ["Entry POI"]),
            "entryDistancePercent": first_value(primary, ["Entry Distance %"]),
            "target": first_value(primary, ["Target"]),
            "targetPoi": first_value(primary, ["Target POI"]),
            "invalidation": first_value(primary, ["Invalidation"]),
            "invalidationPoi": first_value(primary, ["Invalidation POI"]),
            "riskReward": first_value(primary, ["R/R"]),
            "upsidePercent": first_value(primary, ["Target Upside %", "Upside %"]),
            "downsidePercent": first_value(primary, ["Invalidation Down %"]),
            "signalExplanation": first_value(primary, ["Summary Screener", "Section"]),
            "summaryScreener": first_value(primary, ["Summary Screener", "Section"]),
            "signalCount": len(stock_signals),
            "technical": {
                "rsiStatus": first_value(row, ["RSI Status"]),
                "rsi14": first_value(row, ["RSI 14"]),
                "macdPosition": first_value(row, ["Lines Position"]),
                "wavePattern": first_value(row, ["Wave Pattern"]),
                "betaZone": first_value(row, ["Beta (vs IHSG) Zone"]),
                "priceLocation": first_value(row, ["Summary"]),
                "smcSummary": first_value(row, ["Summary"]),
                "marketProfileZone": first_value(row, ["MP Summary"]),
                "maZone": first_value(row, ["MA Zone"]),
                "adrPercent": first_value(row, ["ADR %"]),
                "atrPercent": first_value(row, ["ATR (14) %"]),
                "regime": {
                    "rsZone": first_value(row, ["RS Rating Zone"]),
                    "marketCap": first_value(row, ["Market Cap"]),
                    "marketCapCategory": first_value(row, ["Market Cap Categories"]),
                    "liquidityCategory": first_value(row, ["Liquidity Categories"]),
                    "verdictProfile": first_value(row, ["Verdict Weight Profiles"]),
                },
                "smc": {
                    "strongHigh": first_value(row, ["Strong High"]),
                    "weakHigh": first_value(row, ["Weak High"]),
                    "premiumZone": first_value(row, ["Premium Zone"]),
                    "strongLow": first_value(row, ["Strong Low"]),
                    "weakLow": first_value(row, ["Weak Low"]),
                    "discountZone": first_value(row, ["Discount Zone"]),
                    "equilibrium": first_value(row, ["Equilibrium"]),
                    "closestBullishBlock": first_value(row, ["Closest OB Bull"]),
                    "closestBearishBlock": first_value(row, ["Closest OB Bear"]),
                    "bullishBlockAgeDays": first_value(row, ["OB Bull Age (Days)"]),
                    "bearishBlockAgeDays": first_value(row, ["OB Bear Age (Days)"]),
                    "summary": first_value(row, ["Summary"]),
                },
                "liquidity": {
                    "lot": first_value(row, ["Lot"]),
                    "valueApprox": first_value(row, ["Value (Approx)"]),
                    "averageValue20": first_value(row, ["Average Value 20 D (Approx)"]),
                    "averageVolume20": first_value(row, ["Average Volume 20 D"]),
                    "rvolZone": first_value(row, ["RVOL 20 D Zone"]),
                    "rvolChangeZone": first_value(row, ["RVOL Change Zone"]),
                    "rangeZone": first_value(row, ["ADR & ATR (14) Zone"]),
                },
                "marketProfile": {
                    "ibh": first_value(row, ["IBH"]),
                    "ibl": first_value(row, ["IBL"]),
                    "vsIbl": first_value(row, ["vs IBL"]),
                    "pwh": first_value(row, ["PWH"]),
                    "pwl": first_value(row, ["PWL"]),
                    "vsPwl": first_value(row, ["vs PWL"]),
                    "mdh": first_value(row, ["MDH"]),
                    "mdl": first_value(row, ["MDL"]),
                    "vsMdl": first_value(row, ["vs MDL"]),
                    "summary": first_value(row, ["MP Summary"]),
                },
                "vwapProfiles": {
                    "currentMonth": {
                        "vwap": matching_value(row, contains=("Current MVWAP",), suffix="VWAP"),
                        "zone": matching_value(row, contains=("Current MVWAP",), suffix="VWAP Zone"),
                        "priceSigma": matching_value(row, contains=("Current MVWAP",), suffix="Price σ"),
                        "runningDays": matching_value(row, contains=("Current MVWAP",), suffix="Running Days"),
                    },
                    "previousMonth": {
                        "vwap": matching_value(row, contains=("Previous MVWAP",), suffix="VWAP"),
                        "zone": matching_value(row, contains=("Previous MVWAP",), suffix="VWAP Zone"),
                        "priceSigma": matching_value(row, contains=("Previous MVWAP",), suffix="Price σ"),
                    },
                    "currentQuarter": {
                        "vwap": matching_value(row, contains=("Current QVWAP",), suffix="VWAP"),
                        "zone": matching_value(row, contains=("Current QVWAP",), suffix="VWAP Zone"),
                        "priceSigma": matching_value(row, contains=("Current QVWAP",), suffix="Price σ"),
                    },
                    "previousQuarter": {
                        "vwap": matching_value(row, contains=("Previous QVWAP",), suffix="VWAP"),
                        "zone": matching_value(row, contains=("Previous QVWAP",), suffix="VWAP Zone"),
                        "priceSigma": matching_value(row, contains=("Previous QVWAP",), suffix="Price σ"),
                    },
                    "previousYear": {
                        "vwap": matching_value(row, contains=("Previous Year VWAP",), suffix="VWAP"),
                        "zone": matching_value(row, contains=("Previous Year VWAP",), suffix="VWAP Zone"),
                        "priceSigma": matching_value(row, contains=("Previous Year VWAP",), suffix="Price σ"),
                    },
                },
                "movingAverageDetail": {
                    "ema25DifferencePercent": first_value(row, ["EMA 25 %diff"]),
                    "ema25Position": first_value(row, ["EMA 25 Pos"]),
                    "ema50DifferencePercent": first_value(row, ["EMA 50 %diff"]),
                    "ema50Position": first_value(row, ["EMA 50 Pos"]),
                    "sma200DifferencePercent": first_value(row, ["SMA 200 %diff"]),
                    "sma200Position": first_value(row, ["SMA 200 Pos"]),
                },
                "rsiDetail": {
                    "change1d": first_value(row, ["RSI 14 Δ 1D"]),
                    "average14": first_value(row, ["RSI MA 14"]),
                    "position": first_value(row, ["RSI Pos"]),
                    "cross": first_value(row, ["RSI Cross"]),
                    "divergenceSignal": first_value(row, ["Divergence Signal"]),
                    "divergenceStrength": first_value(row, ["Divergence Strength"]),
                    "divergenceStartDate": first_value(row, ["Divergence Start Date"]),
                    "divergenceConfirmDate": first_value(row, ["Divergence Confirm Date"]),
                },
                "macdDetail": {
                    "signalLine": first_value(row, ["Signal Line"]),
                    "histogram": first_value(row, ["Histogram (EMA3)"]),
                    "cross": first_value(row, ["MACD Cross"]),
                },
                "stochDetail": {
                    "k": first_value(row, ["Stochastic %K"]),
                    "d": first_value(row, ["Stochastic %D"]),
                    "cross": first_value(row, ["Stochastic Cross"]),
                },
                "vwapPosition": next(
                    (
                        value
                        for key, value in row.items()
                        if "VWAP Zone" in key and value not in (None, "", "-")
                    ),
                    None,
                ),
            },
            "fundamentals": {
                "marketCap": first_value(fund, ["Market Cap"]),
                "marketCapClass": first_value(fund, ["Market Cap Categories"]),
                "peRatio": first_value(fund, ["Current PE Ratio (TTM)", "P/E Ratio"]),
                "priceToBook": first_value(fund, ["Current Price to Book Value", "P/B Ratio", "Current PBV"]),
                "roe": first_value(fund, ["Return on Equity (TTM)", "ROE"]),
                "debtToEquity": first_value(fund, ["Debt to Equity Ratio (Quarter)", "Debt to Equity"]),
                "dividendYield": first_value(fund, ["Dividend Yield (%)", "Dividend Yield"]),
                "freeFloat": first_value(fund, ["Free Float (%)", "Free Float %", "Free Float"]),
                "businessSummary": first_value(fund, ["Business Summary"]),
            },
            "news": news_rows.get(ticker, [])[:8],
        }
    return dict(sorted(stocks.items()))


def section_filter(text: str) -> tuple[str, str]:
    match = re.search(r"\bFilter\s+([A-G])\b", text, flags=re.IGNORECASE)
    if match:
        tag = match.group(1).upper()
        return tag, FILTER_LABELS.get(tag, f"Filter {tag}")
    normalized = text.lower()
    for tag, label in FILTER_LABELS.items():
        if label.lower() in normalized:
            return tag, label
    return "", ""


def parse_legacy_screener(ws) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    current_section = ""
    for row_idx in range(1, ws.max_row + 1):
        first = str(ws.cell(row_idx, 2).value or "").strip()
        tag, label = section_filter(first)
        if tag and not re.fullmatch(r"[A-G]", first):
            current_section = first
            continue
        if not re.fullmatch(r"[A-G]", first):
            continue
        values = [clean_value(ws.cell(row_idx, col).value) for col in range(2, 19)]
        item = dict(zip(LEGACY_SCREENER_COLUMNS, values))
        item["Filter Label"] = FILTER_LABELS.get(first, label or f"Filter {first}")
        item["Section"] = current_section or f"Filter {first}"
        rows.append(item)
    return rows


def parse_idx_screener(ws) -> list[dict[str, Any]]:
    header_row = find_header_row(ws, {"Ticker", "Summary Screener"})
    if not header_row:
        return parse_legacy_screener(ws)

    headers = unique_headers(ws, header_row)
    ticker_col = next(col for col, header in headers if header == "Ticker")
    rows: list[dict[str, Any]] = []
    current_filter = ""
    current_label = ""
    current_section = ""

    for row_idx in range(header_row + 1, ws.max_row + 1):
        first = str(ws.cell(row_idx, 2).value or "").strip()
        ticker = str(ws.cell(row_idx, ticker_col).value or "").strip()
        tag, label = section_filter(first)
        if tag and not ticker:
            current_filter, current_label, current_section = tag, label, first
            continue
        if not ticker or ticker.lower().startswith("no tickers"):
            continue

        item = {
            header: clean_value(ws.cell(row_idx, col).value)
            for col, header in headers
        }
        item["Filter"] = current_filter or str(item.get("Filter") or "")
        item["Filter Label"] = current_label or FILTER_LABELS.get(item["Filter"], "")
        item["Section"] = current_section or f"Filter {item['Filter']}"
        rows.append(item)
    return rows


def signal_summary(screener: list[dict[str, Any]]) -> list[dict[str, Any]]:
    counts = Counter(str(row.get("Filter") or "") for row in screener)
    tags = sorted(set(FILTER_LABELS) | set(counts))
    return [
        {"id": tag, "label": FILTER_LABELS.get(tag, f"Filter {tag}"), "count": counts[tag]}
        for tag in tags
        if counts[tag] or tag in FILTER_LABELS
    ]


def market_overview(
    screener: list[dict[str, Any]],
    technical: list[dict[str, Any]],
    processing: list[dict[str, Any]],
) -> dict[str, Any]:
    breadth = {"advances": 0, "declines": 0, "unchanged": 0}
    sectors: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"count": 0, "advances": 0, "declines": 0, "changeTotal": 0.0, "changeCount": 0}
    )
    movers: list[dict[str, Any]] = []

    for row in technical:
        ticker = str(row.get("Ticker") or "").strip()
        change = number(first_value(row, ("Price Change %", "Chg %")))
        price = number(first_value(row, ("Closing Price", "Price")))
        sector = normalize_idx_sector(
            first_value(row, ("IDX Sector",)),
            first_value(row, ("Sector",)),
        )
        if change is not None:
            if change > 0:
                breadth["advances"] += 1
                sectors[sector]["advances"] += 1
            elif change < 0:
                breadth["declines"] += 1
                sectors[sector]["declines"] += 1
            else:
                breadth["unchanged"] += 1
            sectors[sector]["changeTotal"] += change
            sectors[sector]["changeCount"] += 1
            movers.append({"ticker": ticker, "price": price, "change": change, "sector": sector})
        sectors[sector]["count"] += 1

    signal_tickers: dict[str, set[str]] = defaultdict(set)
    for row in screener:
        sector = normalize_idx_sector(
            first_value(row, ("IDX Sector",)),
            first_value(row, ("Sector",)),
        )
        signal_tickers[sector].add(str(row.get("Ticker") or ""))

    sector_rows = []
    for sector, stats in sectors.items():
        count = stats["changeCount"]
        sector_rows.append(
            {
                "sector": sector,
                "count": stats["count"],
                "advances": stats["advances"],
                "declines": stats["declines"],
                "avgChange": stats["changeTotal"] / count if count else None,
                "signals": len(signal_tickers.get(sector, set())),
            }
        )
    sector_rows.sort(key=lambda row: (row["avgChange"] is not None, row["avgChange"] or 0), reverse=True)

    statuses = Counter(str(row.get("Status") or "Unknown") for row in processing)
    movers.sort(key=lambda row: row["change"])
    return {
        "breadth": breadth,
        "signals": signal_summary(screener),
        "sectors": sector_rows,
        "topGainers": list(reversed(movers[-8:])),
        "topDecliners": movers[:8],
        "processingStatus": dict(statuses),
    }


def read_payload(path: Path) -> dict[str, Any] | None:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def rebuild_history() -> None:
    history: dict[str, list[dict[str, Any]]] = defaultdict(list)
    dates: list[str] = []
    for path in sorted(SOURCE_PAYLOAD_DIR.glob("20??-??-??.json")):
        payload = read_payload(path)
        if not payload:
            continue
        market_date = str(payload.get("date") or path.stem)
        dates.append(market_date)
        for row in payload.get("technical", []):
            ticker = str(row.get("Ticker") or "").strip()
            price = number(first_value(row, ("Closing Price", "Price")))
            change = number(first_value(row, ("Price Change %", "Chg %")))
            if ticker and price is not None:
                history[ticker].append(
                    {"date": market_date, "price": price, "change": change}
                )
    payload = {
        "dates": sorted(set(dates)),
        "tickers": {ticker: points for ticker, points in sorted(history.items())},
    }
    (SOURCE_PAYLOAD_DIR / "history.json").write_text(
        json.dumps(payload, separators=(",", ":"), ensure_ascii=False),
        encoding="utf-8",
    )


def update_manifest(entry: dict[str, Any]) -> None:
    manifest_path = DATA_DIR / "manifest.json"
    manifest = read_payload(manifest_path) or {"latest": None, "dates": []}
    if int(manifest.get("schemaVersion") or 0) >= 5:
        # The canonical schema-v5 archive is updated by scripts/archive_v5.py.
        # Avoid mixing legacy `date` entries into its `marketDate` contract.
        return
    dates = [item for item in manifest.get("dates", []) if item.get("date") != entry["date"]]
    dates.append(entry)
    dates.sort(key=lambda item: item["date"], reverse=True)
    manifest["dates"] = dates
    trading_dates = [item for item in dates if item.get("isTradingDate", True)]
    manifest["latest"] = trading_dates[0]["date"] if trading_dates else (dates[0]["date"] if dates else None)
    manifest["history"] = "data/history.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def publish_ohlcv_cache(source: Path, target: Path, max_rows: int = 700) -> None:
    if target.exists():
        shutil.rmtree(target)
    target.mkdir(parents=True, exist_ok=True)
    for source_file in source.glob("*.json"):
        try:
            payload = json.loads(source_file.read_text(encoding="utf-8"))
            payload.setdefault("schemaVersion", 1)
            payload.setdefault("source", "yfinance")
            payload.setdefault("adjusted", False)
            payload.setdefault("timezone", "Asia/Jakarta")
            payload.setdefault("session", "IDX regular daily session")
            payload.setdefault("formulaVersion", "ohlcv-series-v1")
            payload["rows"] = list(payload.get("rows") or [])[-max_rows:]
            for row in payload["rows"]:
                row.setdefault("source", payload["source"])
                row.setdefault("adjusted", payload["adjusted"])
                row.setdefault("timezone", payload["timezone"])
                row.setdefault("session", payload["session"])
            (target / source_file.name).write_text(
                json.dumps(payload, separators=(",", ":"), ensure_ascii=False),
                encoding="utf-8",
            )
        except (OSError, json.JSONDecodeError, TypeError):
            continue


def backfill_stale_fundamentals(fundamental: list[dict[str, Any]], market_date: str) -> int:
    """
    yfinance is rate-limited/blocked hard enough from the CI runner that ~65%
    of tickers come back with an empty `info` dict on a given day, and both
    documented fallback sources are now dead end-to-end (IDX.co.id retired
    the ratio-summary endpoint it used to expose; Investing.com's API sits
    behind Cloudflare bot management). Rather than publish "-" for a ticker
    whose data we HAVE fetched successfully before, backfill it from the most
    recent prior committed date that has real numbers for that ticker — never
    a fabricated value, always something this pipeline itself published on an
    earlier date. Every backfilled row is tagged with "Fundamentals As Of" so
    staleness is disclosed, never silent. Returns the count backfilled.
    """
    dates_dir = DATA_DIR / "dates"
    if not dates_dir.is_dir():
        return 0
    still_needed = {
        str(row.get("Ticker") or "").upper()
        for row in fundamental
        if row.get("Market Cap") in (None, "-", "")
    }
    if not still_needed:
        return 0
    prior_dates = sorted(
        (d.name for d in dates_dir.iterdir() if d.is_dir() and d.name < market_date),
        reverse=True,
    )
    found: dict[str, tuple[dict[str, Any], str]] = {}
    for d in prior_dates:
        if not (still_needed - found.keys()):
            break
        prior_path = dates_dir / d / "fundamental.json"
        if not prior_path.exists():
            continue
        try:
            prior = json.loads(prior_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        for rec in prior.get("records", []):
            ticker = str(rec.get("Ticker") or "").upper()
            if ticker in still_needed and ticker not in found and rec.get("Market Cap") not in (None, "-", ""):
                found[ticker] = (rec, d)

    backfilled = 0
    for row in fundamental:
        ticker = str(row.get("Ticker") or "").upper()
        if ticker not in found or row.get("Market Cap") not in (None, "-", ""):
            continue
        prior_rec, prior_date = found[ticker]
        for key, value in prior_rec.items():
            if key in ("Ticker", "Fundamentals As Of"):
                continue
            if row.get(key) in (None, "-", "") and value not in (None, "-", ""):
                row[key] = value
        row["Fundamentals As Of"] = prior_date
        backfilled += 1
    return backfilled


def export_workbook(workbook_path: Path) -> Path:
    from rebuild_backend.logic_reference import export_registry
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
    SOURCE_PAYLOAD_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(workbook_path, read_only=False, data_only=True)
    if "Data Processing Results" not in wb.sheetnames:
        raise ValueError("Workbook has no Data Processing Results sheet")
    processing_sheet = wb["Data Processing Results"]
    market_date = slug_date(find_stat(processing_sheet, "As-of Date"))

    run_time = clean_value(find_stat(processing_sheet, "Run Time"))
    total_scanned = clean_value(find_stat(processing_sheet, "Total Tickers Scanned"))
    ok_count = clean_value(find_stat(processing_sheet, "OK  (full data)"))
    partial_count = clean_value(find_stat(processing_sheet, "Partial Data"))
    no_data_count = clean_value(find_stat(processing_sheet, "No Data"))

    screener = parse_idx_screener(wb["IDX Screener"])
    technical = parse_table(wb["IDX Technical Detail"])
    fundamental = parse_table(wb["IDX Fundamental Detail"])
    backfilled_count = backfill_stale_fundamentals(fundamental, market_date)
    if backfilled_count:
        print(f"[FUND_BACKFILL] {backfilled_count} tickers backfilled from a prior date's real fundamentals (today's fetch was empty).")
    processing = parse_table(processing_sheet)
    news = parse_table(wb["IDX News"]) if "IDX News" in wb.sheetnames else []
    workbook_sheets = {
        sheet_name: workbook_sheet_payload(wb[sheet_name])
        for sheet_name in wb.sheetnames
    }
    field_catalog = build_field_catalog(workbook_sheets)
    ticker_details = build_ticker_details(workbook_sheets)
    stocks = normalized_stocks(technical, fundamental, news, screener)
    qa_rows = workbook_sheets.get("QA Calculation Audit", {}).get("rows", [])

    overview = market_overview(screener, technical, processing)
    payload = {
        "schemaVersion": 3,
        "date": market_date,
        "runTime": run_time,
        "workbook": f"downloads/{market_date}.xlsx",
        "summary": {
            "totalScanned": total_scanned,
            "ok": ok_count,
            "partial": partial_count,
            "noData": no_data_count,
            "signalRows": len(screener),
            "signalTickers": len({str(row.get("Ticker") or "") for row in screener}),
        },
        "overview": overview,
        "columns": {
            "screener": list(screener[0].keys()) if screener else [],
            "technical": list(technical[0].keys()) if technical else [],
            "fundamental": list(fundamental[0].keys()) if fundamental else [],
            "news": list(news[0].keys()) if news else [],
            "processing": list(processing[0].keys()) if processing else [],
        },
        "screener": screener,
        "technical": technical,
        "fundamental": fundamental,
        "news": news,
        "processing": processing,
        "stocks": stocks,
        "workbookSheets": workbook_sheets,
        "fieldCatalog": field_catalog,
        "tickerDetails": ticker_details,
        "qa": {
            "rows": qa_rows,
            "pass": sum(1 for row in qa_rows if str(row.get("QA Status") or "").upper() == "PASS"),
            "warn": sum(1 for row in qa_rows if str(row.get("QA Status") or "").upper() == "WARN"),
            "fail": sum(1 for row in qa_rows if str(row.get("QA Status") or "").upper() == "FAIL"),
        },
    }

    data_path = SOURCE_PAYLOAD_DIR / f"{market_date}.json"
    data_path.write_text(
        json.dumps(payload, separators=(",", ":"), ensure_ascii=False),
        encoding="utf-8",
    )
    download_path = DOWNLOADS_DIR / f"{market_date}.xlsx"
    if workbook_path.resolve() != download_path.resolve():
        shutil.copy2(workbook_path, download_path)

    qa_dir = SOURCE_PAYLOAD_DIR / "qa"
    qa_dir.mkdir(parents=True, exist_ok=True)
    (qa_dir / f"{market_date}.qa.json").write_text(
        json.dumps(payload["qa"], separators=(",", ":"), ensure_ascii=False),
        encoding="utf-8",
    )

    source_ohlcv = OUTPUT_DIR / "ohlcv" / market_date
    target_ohlcv = DATA_DIR / "ohlcv" / market_date
    if source_ohlcv.exists():
        publish_ohlcv_cache(source_ohlcv, target_ohlcv)

    update_manifest(
        {
            "date": market_date,
            "runTime": run_time,
            "rows": len(screener),
            "tickers": payload["summary"]["signalTickers"],
            "file": f"data_sources/full-workbook/{market_date}.json",
            "workbook": f"downloads/{market_date}.xlsx",
            "qa": f"data_sources/full-workbook/qa/{market_date}.qa.json",
            "ohlcv": f"data/ohlcv/{market_date}",
            "isTradingDate": datetime.strptime(market_date, "%Y-%m-%d").weekday() < 5,
        }
    )
    rebuild_history()
    export_registry(DATA_DIR / "logic-reference.json")
    return data_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=None)
    args = parser.parse_args()

    workbook = args.workbook or latest_workbook()
    data_path = export_workbook(workbook)
    print(f"Exported {workbook.name} -> {data_path}")


if __name__ == "__main__":
    main()
